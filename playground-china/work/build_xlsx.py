#!/usr/bin/env python3
"""Build playground-china/suppliers.xlsx from work/raw/*.jsonl + work/shortlist.json.

Usage: python3 -I playground-china/work/build_xlsx.py
"""
import importlib.util
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "suppliers.xlsx")

sys.dont_write_bytecode = True
spec = importlib.util.spec_from_file_location("merge", os.path.join(HERE, "merge.py"))
merge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(merge)

URL_RE = re.compile(r"https?://[^\s;|,，)）\]]+")
FONT = Font(name="Arial", size=10)
FONT_BOLD = Font(name="Arial", size=10, bold=True)
FONT_LINK = Font(name="Arial", size=10, color="0563C1", underline="single")

CONTACT_LABELS = {
    "contact_person": "лицо",
    "mobile": "моб.",
    "landline": "тел.",
    "wechat": "WeChat",
    "email": "email",
    "qq": "QQ",
    "whatsapp": "WhatsApp",
    "address": "адрес",
}


def urls_in(v):
    if isinstance(v, list):
        out = []
        for x in v:
            out += URL_RE.findall(str(x))
        return list(dict.fromkeys(out))
    return list(dict.fromkeys(URL_RE.findall(str(v or ""))))


def contact_source_urls(m):
    cs = m.get("contact_sources") or {}
    ordered = []
    for k in CONTACT_LABELS:
        for u in cs.get(k, []):
            if u not in ordered:
                ordered.append(u)
    for k, us in cs.items():
        for u in us:
            if u not in ordered:
                ordered.append(u)
    return ordered


def contact_source_map(m):
    cs = m.get("contact_sources") or {}
    urls = contact_source_urls(m)
    if len(urls) <= 1:
        return ""
    lines = []
    for k, us in cs.items():
        label = CONTACT_LABELS.get(k, k)
        idx = ", ".join(f"№{urls.index(u) + 1}" for u in us if u in urls)
        lines.append(f"{label}: {idx}")
    return "\n".join(lines)


def text(v):
    if v is None:
        return ""
    if isinstance(v, list):
        return "\n".join(str(x) for x in v)
    return str(v)


# (header, getter, kind) ; kind: "text" | "url" | "multiurl"
def nth(seq, i):
    return seq[i] if len(seq) > i else ""


COLUMNS = [
    ("Похожесть 1-5", lambda m: m.get("similarity"), "text"),
    ("Название (кит.)", lambda m: m.get("name_cn"), "text"),
    ("Название (англ.)", lambda m: m.get("name_en"), "text"),
    ("Тип", lambda m: m.get("type"), "text"),
    ("Признаки типа", lambda m: m.get("type_evidence"), "text"),
    ("Провинция", lambda m: m.get("province"), "text"),
    ("Город", lambda m: m.get("city"), "text"),
    ("Адрес", lambda m: m.get("address"), "text"),
    ("Сайт", lambda m: m.get("website"), "url"),
    ("1688", lambda m: m.get("url_1688"), "url"),
    ("Alibaba", lambda m: m.get("url_alibaba"), "url"),
    ("Made-in-China", lambda m: m.get("url_mic"), "url"),
    ("Контактное лицо", lambda m: m.get("contact_person"), "text"),
    ("Должность", lambda m: m.get("contact_title"), "text"),
    ("Мобильный", lambda m: m.get("mobile"), "text"),
    ("Стационарный", lambda m: m.get("landline"), "text"),
    ("WeChat", lambda m: m.get("wechat"), "text"),
    ("Email", lambda m: m.get("email"), "text"),
    ("QQ", lambda m: m.get("qq"), "text"),
    ("WhatsApp", lambda m: m.get("whatsapp"), "text"),
    ("Источник контактов №1", lambda m: nth(contact_source_urls(m), 0), "url"),
    ("Источник контактов №2", lambda m: nth(contact_source_urls(m), 1), "url"),
    ("Источник контактов №3", lambda m: nth(contact_source_urls(m), 2), "url"),
    ("Какой контакт откуда", contact_source_map, "text"),
    ("Основная продукция", lambda m: m.get("main_products"), "text"),
    ("Башни от 6 м", lambda m: m.get("towers_6m"), "text"),
    ("Башни: подтверждение", lambda m: m.get("tower_evidence"), "text"),
    ("Пример башни 1", lambda m: nth(urls_in(m.get("tower_examples")), 0), "url"),
    ("Пример башни 2", lambda m: nth(urls_in(m.get("tower_examples")), 1), "url"),
    ("Пример башни 3", lambda m: nth(urls_in(m.get("tower_examples")), 2), "url"),
    ("HPL", lambda m: m.get("hpl"), "text"),
    ("HPL: подтверждение", lambda m: m.get("hpl_evidence"), "multiurl"),
    ("Трубные горки", lambda m: m.get("tube_slides"), "text"),
    ("Горки: подтверждение", lambda m: m.get("tube_evidence"), "multiurl"),
    ("Сертификаты", lambda m: m.get("certs"), "text"),
    ("Экспорт / СНГ", lambda m: m.get("export_cis"), "multiurl"),
    ("Русскоязычный менеджер", lambda m: m.get("russian_manager"), "multiurl"),
    ("Год основания", lambda m: m.get("founded"), "text"),
    ("Сотрудники", lambda m: m.get("employees"), "text"),
    ("Уставный капитал", lambda m: m.get("reg_capital"), "text"),
    ("Комментарий", lambda m: m.get("comment"), "text"),
    ("Все источники", lambda m: m.get("sources"), "multiurl"),
    ("Дата проверки", lambda m: m.get("checked"), "text"),
    ("Где найдено", lambda m: m.get("found_by"), "text"),
    ("Примечания", lambda m: m.get("notes"), "text"),
]

MAX_W = {"url": 45, "multiurl": 60, "text": 60}


def write_cell(ws, r, c, value, kind):
    s = text(value)
    cell = ws.cell(row=r, column=c, value=s if s != "" else None)
    cell.font = FONT
    if isinstance(value, int):
        cell.value = value
    if kind in ("url", "multiurl") and s:
        us = urls_in(s)
        if us:
            cell.hyperlink = us[0]
            cell.font = FONT_LINK
    if "\n" in s or len(s) > 60:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    else:
        cell.alignment = Alignment(vertical="top")
    return s


SPLIT_HEADERS = {"Контактное лицо", "Должность", "Мобильный", "Стационарный", "WeChat", "Email", "QQ", "WhatsApp", "Сайт"}


def display_width(line):
    return sum(2 if ord(ch) >= 0x2E80 else 1 for ch in line)


def write_table(ws, headers_kinds, rows):
    widths = [display_width(h) for h, _ in headers_kinds]
    for c, (h, _) in enumerate(headers_kinds, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = FONT_BOLD
        cell.alignment = Alignment(vertical="top", wrap_text=False)
    for r, row in enumerate(rows, 2):
        for c, ((h, kind), v) in enumerate(zip(headers_kinds, row), 1):
            if h in SPLIT_HEADERS and isinstance(v, str):
                v = re.sub(r"\s*;\s*", "\n", v.strip())
            s = write_cell(ws, r, c, v, kind)
            longest = max((display_width(line) for line in s.split("\n")), default=0)
            widths[c - 1] = max(widths[c - 1], min(longest, MAX_W.get(kind, 60)))
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = max(8, w + 2)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def key_of(m):
    return merge.norm_en(m.get("name_en")) or merge.norm_cn(m.get("name_cn"))


def main():
    merged = merge.build()

    shortlist_path = os.path.join(HERE, "shortlist.json")
    shortlist = []
    if os.path.exists(shortlist_path):
        with open(shortlist_path, encoding="utf-8") as f:
            shortlist = json.load(f)

    wb = Workbook()
    wb._named_styles["Normal"].font = Font(name="Arial", size=10)

    # Шортлист
    ws = wb.active
    ws.title = "Шортлист"
    by_key = {}
    for m in merged:
        by_key[merge.norm_en(m.get("name_en"))] = m
        by_key[merge.norm_cn(m.get("name_cn"))] = m
    hk = [("Ранг", "text"), ("Почему в шортлисте", "text")] + [(h, k) for h, _, k in COLUMNS]
    rows = []
    for i, item in enumerate(shortlist, 1):
        m = by_key.get(merge.norm_en(item.get("name_en"))) or by_key.get(merge.norm_cn(item.get("name_cn")))
        if m is None:
            raise SystemExit(f"shortlist entry not found: {item}")
        rows.append([i, item.get("reason", "")] + [g(m) for _, g, _ in COLUMNS])
    write_table(ws, hk, rows)

    # Все
    ws = wb.create_sheet("Все")
    hk = [("№", "text")] + [(h, k) for h, _, k in COLUMNS]
    rows = [[i] + [g(m) for _, g, _ in COLUMNS] for i, m in enumerate(merged, 1)]
    write_table(ws, hk, rows)

    # Поиск
    ws = wb.create_sheet("Поиск")
    logs = []
    raw = os.path.join(HERE, "raw")
    for fn in sorted(os.listdir(raw)):
        if fn.endswith("_searchlog.jsonl"):
            with open(os.path.join(raw, fn), encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if line:
                        try:
                            logs.append(json.loads(line))
                        except json.JSONDecodeError:
                            pass
    order = {"lead": 0}
    logs.sort(key=lambda x: (order.get(x.get("agent"), 1), x.get("agent", ""), x.get("source", "")))
    hk = [("Агент", "text"), ("Источник", "text"), ("Запрос", "text"), ("URL", "url"), ("Результат", "text"), ("Дата", "text")]
    rows = [[x.get("agent"), x.get("source"), x.get("query"), x.get("url"), x.get("result"), x.get("date")] for x in logs]
    write_table(ws, hk, rows)

    wb.save(OUT)
    print(f"companies: {len(merged)}; shortlist: {len(shortlist)}; search log rows: {len(logs)} -> {OUT}")


if __name__ == "__main__":
    main()
